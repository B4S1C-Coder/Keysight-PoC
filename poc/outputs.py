from pathlib import Path
from openpyxl import Workbook
from openpyxl.styles import Font

from .models import ExtractionResult

HEADERS = [
    "Attribute",
    "Value",
    "Unit",
    "Tolerance",
    "Note",
    "Source",
]

def write_xlsx(result: ExtractionResult, out_dir: Path = Path.home() / "Documents") -> Path:
    out_dir.mkdir(parents=True, exist_ok=True)

    filename = f"Attributes_{result.mpn}_{result.manufacturer}.xlsx"
    path = out_dir / filename

    wb = Workbook()

    for title, rows in [
        ("Attribute Set-1", result.set1),
        ("Attribute Set-2", result.set2),
    ]:
        ws = wb.active if title == "Attribute Set-1" else wb.create_sheet()

        if not ws:
            raise RuntimeError("Failed to create sheet.")

        ws.title = title
        ws.append(HEADERS)

        for cell in ws[1]:
            cell.font = Font(bold=True)

        for row in rows:
            ws.append([
                row.attribute,
                row.value,
                row.unit,
                row.tolerance,
                row.note,
                row.source,
            ])

        ws.freeze_panes = "A2"
        ws.column_dimensions["A"].width = 22
        ws.column_dimensions["B"].width = 32

    wb.save(path)
    return path