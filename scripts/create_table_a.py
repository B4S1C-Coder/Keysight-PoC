
from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill


OUTPUT_FILE = Path(__file__).parent.parent / "data" / "Table_A.xlsx"

MPNS = [
    "91864A028",
    "90044A249",
    "91251A107",
    "91251A105",
]

HEADERS = ["MPN", "Manufacturer", "Material", "RoHS Compliance"]


def main() -> None:
    OUTPUT_FILE.parent.mkdir(parents=True, exist_ok=True)

    workbook = Workbook()
    sheet = workbook.active

    if not sheet:
        raise RuntimeError("Unable to get sheet.")

    sheet.title = "Table A"
    sheet.append(HEADERS)

    for mpn in MPNS:
        sheet.append([mpn, "MCMASTER CARR SUPPLY", "", ""])

    # for cell in sheet[1]:
    #     cell.font = Font(bold=True, color="FFFFFF")
    #     cell.fill = PatternFill(
    #         "solid", fgColor="1F4E78"
    #     )

    for column, width in {
        "A": 16, "B": 28, "C": 32, "D": 42
    }.items():
        sheet.column_dimensions[column].width = width

    sheet.freeze_panes = "A2"
    workbook.save(OUTPUT_FILE)
    print(f"Created: {OUTPUT_FILE}")


if __name__ == "__main__":
    main()
