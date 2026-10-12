from pathlib import Path
from openpyxl import load_workbook
from playwright.sync_api import sync_playwright
from poc.mcmaster import fetch_attributes

EXCEL_FILE = Path(__file__).parent / "data" / "Table_A.xlsx"

def main() -> None:
    if not EXCEL_FILE.exists():
        raise FileNotFoundError(
            f"Table A not found: {EXCEL_FILE}\n"
            "Run: uv run python poc/create_table_a.py"
        )

    workbook = load_workbook(EXCEL_FILE)
    sheet = workbook["Table A"]

    headers = {
        cell.value: index
        for index, cell in enumerate(sheet[1], start=1)
    }

    mpn_col = headers["MPN"]
    material_col = headers["Material"]
    rohs_col = headers["RoHS Compliance"]

    allowed_mpns = {
        str(sheet.cell(row, mpn_col).value).strip(): row
        for row in range(2, sheet.max_row + 1)
        if sheet.cell(row, mpn_col).value
    }

    mpn = input("Enter McMaster MPN: ").strip()

    if mpn not in allowed_mpns:
        print("Invalid MPN. Allowed MPNs:")
        print("\n".join(sorted(allowed_mpns)))
        workbook.close()
        return

    with sync_playwright() as playwright:
        browser = playwright.chromium.launch(headless=True)

        try:
            page = browser.new_page()
            attributes = fetch_attributes(page, mpn)

        except RuntimeError as err:
            print(f"Unable to fetch {mpn}.\nError:{err}")
        
        finally:
            browser.close()

    row = allowed_mpns[mpn]
    sheet.cell(row, material_col).value = attributes["Material"]
    sheet.cell(row, rohs_col).value = attributes["RoHS Compliance"]

    workbook.save(EXCEL_FILE)
    workbook.close()

    print(f"\nUpdated {EXCEL_FILE}")
    print(f"MPN: {mpn}")
    print(f"Material: {attributes['Material']}")
    print(f"RoHS Compliance: {attributes['RoHS Compliance']}")

if __name__ == "__main__":
    main()